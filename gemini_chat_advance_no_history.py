import os
import re
import json
import docx
import pandas as pd
import configparser
import keyring
import datetime
from rich import print as rprint
from rich.console import Console
from rich.syntax import Syntax
from rich.markdown import Markdown
from rich.panel import Panel
import google.generativeai as genai
from PyPDF2 import PdfReader
from cryptography.fernet import Fernet
from email.parser import Parser
from email.policy import default
import glob
import pyperclip
import keyboard
import argparse
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
import openpyxl

console = Console()
all_code_blocks = []

# --- (API Key Management - unchanged) ---
def encrypt_api_key(api_key, key):
    f = Fernet(key)
    return f.encrypt(api_key.encode()).decode()

def decrypt_api_key(encrypted_api_key, key):
    f = Fernet(key)
    return f.decrypt(encrypted_api_key.encode()).decode()

def get_api_key():
    config = configparser.ConfigParser()
    if not os.path.exists('config.ini'):
        api_key = console.input("[bold yellow]Enter your Gemini API key: [/bold yellow]")
        encryption_key = Fernet.generate_key()
        keyring.set_password("gemini_chat", "encryption_key", encryption_key.decode())
        encrypted_api_key = encrypt_api_key(api_key, encryption_key)
        config['API'] = {'encrypted_key': encrypted_api_key}
        with open('config.ini', 'w') as configfile:
            config.write(configfile)
    else:
        config.read('config.ini')

    encrypted_api_key = config['API']['encrypted_key']
    encryption_key = keyring.get_password("gemini_chat", "encryption_key").encode()
    return decrypt_api_key(encrypted_api_key, encryption_key)

api_key = get_api_key()
genai.configure(api_key=api_key)


# --- File Handling Functions (Improved Regex) ---

def detect_file_in_prompt(prompt):
    """Detects a file path, converts it to an absolute path, and checks if it exists."""
    match = re.search(r'([a-zA-Z]:[\\/](?:[\w\s\-\._()]+[\\/]?)+)', prompt)
    if match:
        file_path = match.group(0)
        # Convert to absolute path *immediately*
        absolute_path = os.path.abspath(file_path)

        if os.path.exists(absolute_path):
            return absolute_path  # Return the absolute path
        else:
            rprint(f"[bold yellow]Detected path '{file_path}' (absolute: '{absolute_path}') does not exist.[/bold yellow]")
            return None
    else:
        return None

def summarize_text(text, max_length=5000):
    """Summarizes long text using Gemini."""
    if len(text) <= max_length:
        return text
    summary_prompt = f"Summarize the following text:\n{text[:max_length]}..."
    # Use the modified generate_response with retries:
    return generate_response(summary_prompt)

def extract_from_json(file_path, ai_suggestion=None):
    """Extracts data from a JSON file, optionally using AI suggestions."""
    try:
        with open(file_path, 'r', encoding='utf-8') as file:
            data = json.load(file)
            if isinstance(data, dict) and ai_suggestion:
                keys_to_extract = [key.strip() for key in ai_suggestion.split(",")]
                return {key: data.get(key, "N/A") for key in keys_to_extract}
            return json.dumps(data, indent=2)  # Return pretty-printed JSON
    except Exception as e:
        return f"[Error reading JSON: {str(e)}]"

def extract_from_excel(file_path, ai_suggestion=None):
    """Extracts data from Excel files, handling sheets and columns."""
    try:
        df = pd.read_excel(file_path, sheet_name=None)  # Read all sheets
        if ai_suggestion and ":" in ai_suggestion:
            sheet_name, column_name = ai_suggestion.split(":", 1)  # Split only once
            sheet_name = sheet_name.strip()
            column_name = column_name.strip()

            if sheet_name in df and column_name in df[sheet_name].columns:
                return df[sheet_name][column_name].to_string(index=False)
            return f"[Sheet '{sheet_name}' or Column '{column_name}' not found]"

        # If no specific sheet/column, return the first sheet's content
        first_sheet_name = list(df.keys())[0]
        return df[first_sheet_name].to_string(index=False)
    except Exception as e:
        return f"[Error reading Excel: {str(e)}]"

def read_file_content(file_path, ai_suggestion=None):
    """Reads content from various file types, using AI suggestions where applicable."""
    if not os.path.exists(file_path):
        return "[Error] File not found."

    _, ext = os.path.splitext(file_path)
    try:
        if ext in ['.txt', '.py', '.cs', '.js']:
            with open(file_path, 'r', encoding='utf-8') as file:
                return summarize_text(file.read())
        elif ext == '.json':
            return extract_from_json(file_path, ai_suggestion)
        elif ext in ['.xlsx', '.xls']:
            return extract_from_excel(file_path, ai_suggestion)
        elif ext == '.csv':
            df = pd.read_csv(file_path)
            return df.head().to_string(index=False)  # Show first few rows
        elif ext == '.log':
            with open(file_path, 'r', encoding='utf-8') as file:
                lines = file.readlines()
                errors = [line for line in lines if "ERROR" in line or "WARN" in line]
                return "\n".join(errors[:10]) if errors else "[No errors/warnings found]"
        elif ext == '.pdf':
            reader = PdfReader(file_path)
            content = "\n".join(page.extract_text() for page in reader.pages if page.extract_text())
            return summarize_text(content)
        elif ext == '.docx':  # Handle .docx files
            doc = docx.Document(file_path)
            full_text = []
            for paragraph in doc.paragraphs:
                full_text.append(paragraph.text)
            return summarize_text('\n'.join(full_text))
        else:
            return f"[Unsupported file type: {ext}]"
    except Exception as e:
        return f"[Error reading file: {str(e)}]"
    
# --- Core Gemini Interaction (Key Changes Here) ---

def generate_response(prompt, model_name="gemini-2.0-flash", max_retries=5, retry_delay=5):
    """Generates a response, handling rate limits with retries."""
    model = genai.GenerativeModel(model_name)
    for attempt in range(max_retries):
        try:
            response = model.generate_content(prompt)
            return response.text
        except genai.APIError as e:
            if e.status_code == 429:  # Rate limit error
                rprint(f"[bold yellow]Rate limit exceeded (attempt {attempt + 1}/{max_retries}). Waiting {retry_delay} seconds...[/bold yellow]")
                datetime.time.sleep(retry_delay)
                retry_delay *= 2  # Exponential backoff (increase delay each retry)
            else:
                rprint(f"[bold red]An API error occurred: {e}[/bold red]")
                return f"[Error] An API error occurred: {e}"  # Return error message
        except Exception as e: #catch other exception
            rprint(f"[bold red]An error occurred: {e}[/bold red]")
            return f"[Error] An error occurred: {e}"
    rprint(f"[bold red]Max retries exceeded. Failed to generate response.[/bold red]")
    return "[Error] Max retries exceeded"  # Or raise the exception, depending on your needs

def process_prompt_with_file(prompt):
    """Processes prompts that include file paths."""
    file_path = detect_file_in_prompt(prompt)
    if not file_path:
        return generate_response(prompt)  # No file, normal generation

    # 1. ISOLATE the file path in the prompt.  This is VERY important.
    #    We replace the file path with a placeholder.
    placeholder = "<<<FILE_PATH>>>"
    prompt_with_placeholder = prompt.replace(file_path, placeholder)

    # 2.  More focused intent analysis.
    intent_prompt = (
        f"The user mentions a file in their prompt.  Here's the prompt, "
        f"with the file path replaced by '{placeholder}':\n\n"
        f"'{prompt_with_placeholder}'\n\n"
        f"What specific information or action is the user likely requesting "
        f"regarding the file?  Examples: 'summarize', 'extract key points', "
        f"'find errors', 'get the value of X', 'list column names'.  "
        f"Be concise and specific."
    )

    ai_intent = generate_response(intent_prompt)
    rprint(f"[bold blue]AI Intent Analysis:[/bold blue] {ai_intent}")

    # 3. Read the file content.
    file_content = read_file_content(file_path, ai_intent)

    # 4. Construct the final prompt.  Notice how we use the placeholder.
    final_prompt = (
        f"User's request (with file path replaced by '{placeholder}'):\n"
        f"{prompt_with_placeholder}\n\n"
        f"Actual file path: {file_path}\n\n"  # Include the *actual* path separately
        f"Extracted File Content (or error):\n{file_content}\n\n"
        f"Based on the file content and AI's interpretation of the user's "
        f"intent ({ai_intent}), please fulfill the user's request."
    )
    return generate_response(final_prompt)

def display_formatted_response(response_text):
    """Identifies and formats code and markdown in the response (from original)."""
    global all_code_blocks
    segments = response_text.split("```")
    for i, segment in enumerate(segments):
        if i % 2 == 0:  # Markdown section
            console.print(Markdown(segment.strip()))
        else:  # Code block
            lines = segment.split('\n')
            language = lines[0].strip() if lines[0].strip() else 'python'
            code = '\n'.join(lines[1:])
            syntax = Syntax(code, language, theme="monokai", line_numbers=True)
            console.print(Panel(syntax, title=f"[bold]{language.capitalize()} Code[/bold]", expand=True))
            all_code_blocks.append(code)

def get_user_prompt():
    """Prompts the user for input (from original)."""
    try:
        prompt = console.input("[bold green]Enter your prompt (or 'quit'/'exit' to terminate): [/bold green]")
        if prompt.lower() in ['quit', 'exit']:
            quit_application()
        return prompt
    except KeyboardInterrupt:
        quit_application()

def copy_to_clipboard():
    """Copies all generated code blocks to clipboard (from original)."""
    global all_code_blocks
    if all_code_blocks:
        all_code = "\n\n".join(all_code_blocks)
        pyperclip.copy(all_code)
        rprint("\n[bold green]All code blocks copied to clipboard![/bold green]")
    else:
        rprint("\n[bold yellow]No code available to copy.[/bold yellow]")

def quit_application(signum=None, frame=None):
    """Gracefully exits the application (from original)."""
    rprint("\n[bold red]Exiting the program. Goodbye![/bold red]")
    keyboard.unhook_all()  # Unhook hotkeys
    exit(0)



def ensure_directory_exists(file_path):
    """Ensure the directory for the given file path exists (from original)."""
    directory = os.path.dirname(file_path)
    if directory and not os.path.exists(directory):
        try:
            os.makedirs(directory)
            rprint(f"[bold blue]Created directory: {directory}[/bold blue]")
        except Exception as e:
            rprint(f"[bold red]Error creating directory {directory}: {str(e)}[/bold red]")
            return False  # Indicate failure
    return True




# --- Main Function (combined and adapted) ---

def main():
    parser = argparse.ArgumentParser(description='Enhanced Gemini Chat with File Processing')
    subparsers = parser.add_subparsers(dest='command', help='Commands')

    # Chat mode (original functionality)
    chat_parser = subparsers.add_parser('chat', help='Start interactive chat mode')

    args = parser.parse_args()

    keyboard.add_hotkey('ctrl+alt+c', copy_to_clipboard)  # Hotkey setup

       
    # Default to chat mode
    rprint("[bold]Welcome to the Enhanced Gemini Chat![/bold]")
    rprint("[bold cyan]Press Ctrl+Alt+C to copy all generated code blocks.[/bold cyan]")
    rprint("[bold cyan]Press Ctrl+C or type 'quit'/'exit' to terminate the program.[/bold cyan]")

    while True:
        try:
            prompt = get_user_prompt()
            if not prompt:
                continue

            rprint("[bold yellow]Generating response...[/bold yellow]")
            response = process_prompt_with_file(prompt) # Use the combined processing function
            display_formatted_response(response)

        except Exception as e:
            console.print(f"[bold red]An error occurred: {str(e)}[/bold red]")

if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        quit_application()